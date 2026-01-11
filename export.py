"""Export module for sermon documents.

Provides functions to export sermons to Word (.docx) and PDF formats.
"""
import io
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export_to_word(sermon):
    """Export a sermon to Word document format.

    Args:
        sermon: Dictionary containing sermon data with keys:
            - title (required): Sermon title
            - scripture (required): Scripture reference
            - manuscript (optional): Full sermon text
            - outline (optional): Sermon outline

    Returns:
        bytes: The Word document content as bytes
    """
    doc = Document()

    # Add title
    title = doc.add_heading(sermon.get('title', 'Untitled Sermon'), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add scripture reference
    scripture_para = doc.add_paragraph()
    scripture_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scripture_run = scripture_para.add_run(sermon.get('scripture', ''))
    scripture_run.italic = True

    # Add outline if available
    outline = sermon.get('outline')
    if outline:
        doc.add_heading('Outline', level=1)
        for line in outline.split('\n'):
            if line.strip():
                doc.add_paragraph(line.strip())

    # Add manuscript
    manuscript = sermon.get('manuscript')
    if manuscript:
        doc.add_heading('Manuscript', level=1)
        # Split into paragraphs and add each
        paragraphs = manuscript.split('\n\n')
        for para_text in paragraphs:
            if para_text.strip():
                doc.add_paragraph(para_text.strip())

    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def export_to_pdf(sermon):
    """Export a sermon to PDF format.

    Args:
        sermon: Dictionary containing sermon data with keys:
            - title (required): Sermon title
            - scripture (required): Scripture reference
            - manuscript (optional): Full sermon text
            - outline (optional): Sermon outline

    Returns:
        bytes: The PDF document content as bytes
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch

    title = sermon.get('title', 'Untitled Sermon')
    scripture = sermon.get('scripture', '')
    manuscript = sermon.get('manuscript', '')
    outline = sermon.get('outline', '')

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter,
                            leftMargin=inch, rightMargin=inch,
                            topMargin=inch, bottomMargin=inch,
                            title=title,
                            author='AutomatedPastor')

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        alignment=TA_CENTER,
        spaceAfter=12
    )

    scripture_style = ParagraphStyle(
        'Scripture',
        parent=styles['Normal'],
        alignment=TA_CENTER,
        fontName='Times-Italic',
        spaceAfter=24
    )

    heading_style = ParagraphStyle(
        'Heading',
        parent=styles['Heading2'],
        fontSize=14,
        spaceAfter=12
    )

    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=12,
        leading=18,
        firstLineIndent=36,
        spaceAfter=6
    )

    story = []

    # Add title
    story.append(Paragraph(title, title_style))

    # Add scripture
    story.append(Paragraph(scripture, scripture_style))

    # Add outline if available
    if outline:
        story.append(Paragraph('Outline', heading_style))
        for line in outline.split('\n'):
            if line.strip():
                story.append(Paragraph(line.strip(), body_style))
        story.append(Spacer(1, 12))

    # Add manuscript
    if manuscript:
        story.append(Paragraph('Manuscript', heading_style))
        paragraphs = manuscript.split('\n\n')
        for para in paragraphs:
            if para.strip():
                story.append(Paragraph(para.strip(), body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.read()
